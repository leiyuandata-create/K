#!/usr/bin/env python3
"""export_pptx.py <lessondir>: slides.pdf (slide | notes) + slides.tex -> slides.pptx  (spec 9.2)

* Each slide is one 3840 x 2160 PNG of the left half of the joined PDF, full bleed against the
  presentation's own sldSz.  Nothing else is placed on the slide.
* Notes come from the \\note{} bodies in slides.tex (brace matching), split at \\par.
  $...$ becomes a native PowerPoint equation (pandoc/texmath -> OMML inside mc:AlternateContent /
  a14:m, Cambria Math a:rPr on every math run) with a linear fallback.  One pandoc call for all
  equations; the returned count is asserted.
* Alignment: frames = PDF pages = PPTX slides, or stop without writing.
"""
import copy, os, re, subprocess, sys, tempfile, zipfile
import pymupdf as fitz
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
    'a14': 'http://schemas.microsoft.com/office/drawing/2010/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
}
A = '{%s}' % NS['a']; M = '{%s}' % NS['m']; MC = '{%s}' % NS['mc']; A14 = '{%s}' % NS['a14']
NOTE_SZ = '1600'

# project macros that may appear inside notes, expanded before conversion
MACROS = [
    (r'\dydx', r'\dfrac{dy}{dx}'),
    (r'\mps', r'\,\mathrm{m\,s^{-1}}'),
    (r'\mpss', r'\,\mathrm{m\,s^{-2}}'),
]


def matching_brace(s, i):
    """s[i] == '{'; index of its matching '}'."""
    depth = 0
    for j in range(i, len(s)):
        if s[j] == '\\':
            continue
        if s[j] == '{' and (j == 0 or s[j - 1] != '\\'):
            depth += 1
        elif s[j] == '}' and s[j - 1] != '\\':
            depth -= 1
            if depth == 0:
                return j
    raise ValueError('unbalanced braces')


def frames_and_notes(tex):
    body = tex[tex.index('\\begin{document}'):]
    frames = re.split(r'\\begin\{frame\}', body)[1:]
    notes = []
    for fr in frames:
        fr = fr[:fr.index('\\end{frame}')]
        k = fr.find('\\note{')
        if k < 0:
            notes.append('')
            continue
        j = matching_brace(fr, k + 5)
        notes.append(fr[k + 6:j])
    return notes


def expand(s):
    for a, b in MACROS:
        s = re.sub(re.escape(a) + r'(?![A-Za-z])', lambda m: b, s)
    s = re.sub(r'\\unit\{([^{}]*)\}', r'\\,\\mathrm{\1}', s)
    return s


def prose(s):
    s = s.replace('\\%', '%').replace('~', ' ').replace('\\textperiodcentered', '\u00b7')
    s = s.replace('---', '\u2014').replace('--', '\u2013').replace("``", '\u201c').replace("''", '\u201d')
    s = s.replace('\\ ', ' ').replace('\\,', ' ').replace('\\&', '&')
    return re.sub(r'\s+', ' ', s)


def linear(tex):
    """Readable linear fallback for software without equation support."""
    s = tex
    for _ in range(4):
        s = re.sub(r'\\[dt]?frac\{([^{}]*)\}\{([^{}]*)\}', r'(\1)/(\2)', s)
    s = re.sub(r'\^\{([^{}]*)\}', r'^(\1)', s)
    s = re.sub(r'_\{([^{}]*)\}', r'_\1', s)
    rep = {'\\times': '\u00d7', '\\to': '\u2192', '\\pm': '\u00b1', '\\le': '\u2264', '\\ge': '\u2265',
           '\\ne': '\u2260', '\\cdot': '\u00b7', '\\pi': '\u03c0', '\\sqrt': '\u221a', '\\approx': '\u2248',
           '\\Bigl': '', '\\Bigr': '', '\\left': '', '\\right': '', '\\mathrm': '', '\\,': ' ',
           '\\quad': '  ', '\\;': ' ', '\\infty': '\u221e', '\\text': '', '\\dots': '\u2026', '\\ldots': '\u2026'}
    for k, v in rep.items():
        s = s.replace(k, v)
    s = s.replace('{', '').replace('}', '').replace('-', '\u2212')
    return s


def split_segments(par):
    """[(kind, text)] with kind 'text' or 'math'."""
    out, pos = [], 0
    for m in re.finditer(r'\$([^$]+)\$', par):
        if m.start() > pos:
            out.append(('text', par[pos:m.start()]))
        out.append(('math', m.group(1)))
        pos = m.end()
    if pos < len(par):
        out.append(('text', par[pos:]))
    return out


def omml_all(maths):
    """One pandoc call; returns one m:oMath element per input, asserted."""
    if not maths:
        return []
    with tempfile.TemporaryDirectory() as d:
        md = os.path.join(d, 'eq.md'); dx = os.path.join(d, 'eq.docx')
        with open(md, 'w') as f:
            for i, t in enumerate(maths):
                f.write('EQ%d $%s$\n\n' % (i, t))
        subprocess.run(['pandoc', '-f', 'markdown', '-t', 'docx', md, '-o', dx], check=True)
        xml = zipfile.ZipFile(dx).read('word/document.xml')
    root = etree.fromstring(xml)
    paras = root.findall('.//w:body/w:p', NS)
    res = []
    for pp in paras:
        om = pp.findall('.//m:oMath', NS)
        txt = ''.join(pp.xpath('.//w:t/text()', namespaces=NS)).strip()
        if txt.startswith('EQ'):
            if len(om) != 1:
                raise SystemExit('pandoc returned %d equations for %r' % (len(om), txt))
            res.append(om[0])
    if len(res) != len(maths):
        raise SystemExit('pandoc returned %d equations for %d inputs' % (len(res), len(maths)))
    return res


def math_rpr(italic):
    r = etree.Element(A + 'rPr', lang='en-US', sz=NOTE_SZ, i='1' if italic else '0')
    etree.SubElement(r, A + 'latin', typeface='Cambria Math', panose='02040503050406030204',
                     pitchFamily='18', charset='0')
    etree.SubElement(r, A + 'cs', typeface='Cambria Math', panose='02040503050406030204',
                     pitchFamily='18', charset='0')
    return r


def to_ppt_math(om):
    om = copy.deepcopy(om)
    for r in om.iter(M + 'r'):
        rpr = r.find(M + 'rPr')
        plain = rpr is not None and rpr.find(M + 'sty') is not None
        pos = 1 if rpr is not None else 0
        r.insert(pos, math_rpr(not plain))
    for c in om.iter(M + 'ctrlPr'):
        for ch in list(c):
            c.remove(ch)
        c.append(math_rpr(True))
    return om


def text_run(txt):
    r = etree.Element(A + 'r')
    etree.SubElement(r, A + 'rPr', lang='en-US', sz=NOTE_SZ, dirty='0')
    t = etree.SubElement(r, A + 't'); t.text = txt
    return r


def build_paragraph(segs, maths_iter):
    p = etree.Element(A + 'p')
    for kind, txt in segs:
        if kind == 'text':
            t = prose(txt)
            if t:
                p.append(text_run(t))
        else:
            src, om = next(maths_iter)
            ac = etree.SubElement(p, MC + 'AlternateContent', nsmap={'mc': NS['mc']})
            ch = etree.SubElement(ac, MC + 'Choice', Requires='a14', nsmap={'a14': NS['a14']})
            m = etree.SubElement(ch, A14 + 'm')
            m.append(to_ppt_math(om))
            fb = etree.SubElement(ac, MC + 'Fallback')
            fb.append(text_run(linear(src)))
    end = etree.SubElement(p, A + 'endParaRPr', lang='en-US', sz=NOTE_SZ, dirty='0')
    return p


def main(lesson):
    here = os.path.dirname(os.path.abspath(__file__))
    ld = os.path.join(here, lesson)
    tex = open(os.path.join(ld, 'slides.tex')).read()
    nframes = len(re.findall(r'\\begin\{frame\}', tex[tex.index('\\begin{document}'):]))
    notes = frames_and_notes(tex)
    doc = fitz.open(os.path.join(ld, 'slides.pdf'))
    if not (nframes == len(notes) == doc.page_count):
        raise SystemExit('alignment: %d frames, %d notes, %d PDF pages -- not writing' % (nframes, len(notes), doc.page_count))

    # notes -> paragraphs -> segments; collect maths for one pandoc call
    all_pars, maths = [], []
    for n in notes:
        pars = [p.strip() for p in re.split(r'\\par\b', expand(n))]
        pars = [p for p in pars if p]
        segs = [split_segments(p) for p in pars]
        for sg in segs:
            maths += [t for k, t in sg if k == 'math']
        all_pars.append(segs)
    omml = omml_all(maths)
    it = iter(zip(maths, omml))

    prs = Presentation()
    prs.slide_width = Emu(12192000); prs.slide_height = Emu(6858000)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]
    with tempfile.TemporaryDirectory() as d:
        for i, page in enumerate(doc):
            r = page.rect
            clip = fitz.Rect(0, 0, r.width / 2, r.height)
            mat = fitz.Matrix(3840 / clip.width, 2160 / clip.height)
            pix = page.get_pixmap(matrix=mat, clip=clip, alpha=False)
            if (pix.width, pix.height) != (3840, 2160):
                raise SystemExit('slide %d rendered %dx%d' % (i + 1, pix.width, pix.height))
            png = os.path.join(d, 's%03d.png' % i); pix.save(png)
            s = prs.slides.add_slide(blank)
            for sh in list(s.shapes):
                sh._element.getparent().remove(sh._element)
            s.shapes.add_picture(png, 0, 0, width=W, height=H)
            tf = s.notes_slide.notes_text_frame
            txBody = tf._txBody
            for p in txBody.findall(A + 'p'):
                txBody.remove(p)
            bp = txBody.find(A + 'bodyPr')
            for ch in list(bp):
                bp.remove(ch)
            etree.SubElement(bp, A + 'noAutofit')
            for segs in all_pars[i]:
                txBody.append(build_paragraph(segs, it))
            if not all_pars[i]:
                txBody.append(build_paragraph([], it))
        out = os.path.join(ld, 'slides.pptx')
        prs.save(out)
    left = sum(1 for _ in it)
    if left:
        raise SystemExit('%d equations not placed' % left)
    print('%s: slides.pptx written, %d slides, %d native equations' % (lesson, len(prs.slides), len(maths)))


if __name__ == '__main__':
    main(sys.argv[1])
